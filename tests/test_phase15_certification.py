"""Phase 15: run observability, circuit recovery, and output failure tests."""

from __future__ import annotations

import hashlib
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from src.agents.output_agent import run_output
from src.agents.role_agent import run_role_classification
from src.graph.pipeline import build_graph
from src.llm.base import NullLLMProvider
from src.models.job import DateSource, EmploymentType, RawJobPosting
from src.models.state import PipelineState, SourceHealth
from src.services.discovery_orchestrator import orchestrate
from src.services.freshness import is_fresh
from src.services.production_report import production_source_state, render_pipeline_health, source_completeness
from src.services.rejection_codes import rejection_code
from src.services.xlsx import ALL_COLUMNS, iter_workbook_rows, write_workbooks
from src.sources.workday import cxs_request_body
from src.utils.dates import utcnow
from tests.conftest import make_job
from tests.test_discovery_orchestrator import _posting, _profile
from tests.test_llm_reliability import ScriptLLM, _settings

WORK = "Write production code and run code review for backend services."
ROOT = Path(__file__).resolve().parents[1]


def _job(**overrides):
    data = dict(
        job_title="Software Engineer",
        location="Boston, MA",
        description="Required Qualifications\n0-2 years of experience.\n" + WORK,
        posted_at=utcnow() - timedelta(hours=2),
        date_source=DateSource.POSTED_DATE,
        employment_type=EmploymentType.FULL_TIME,
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/swe",
        job_id="swe",
    )
    data.update(overrides)
    return make_job(**data)


def _role_json(*, accept: bool, confidence: float) -> str:
    import json

    return json.dumps(
        {
            "is_software_engineering": accept,
            "role_family": "MACHINE_LEARNING_ENGINEER" if accept else "NOT_SOFTWARE",
            "confidence": confidence,
            "reasoning": "fixture decision",
        }
    )


def test_graph_order_is_unchanged() -> None:
    edges = {(edge.source, edge.target) for edge in build_graph().get_graph().edges}
    assert ("discovery", "extraction") in edges
    assert ("seniority", "location") in edges
    assert ("location", "freshness") in edges
    assert ("freshness", "url") in edges
    assert ("qc", "intelligence") in edges


def test_freshness_boundary_and_workday_body_are_unchanged() -> None:
    now = datetime(2026, 9, 27, 12, tzinfo=timezone.utc)
    posted = RawJobPosting(
        source="greenhouse",
        company_name="Example",
        title="Software Engineer",
        posted_at=now - timedelta(hours=24),
        date_source=DateSource.POSTED_DATE,
        discovered_at=now,
    )
    assert is_fresh(posted, 24, now=now)[0] is True
    assert is_fresh(
        posted.model_copy(update={"posted_at": now - timedelta(hours=24, seconds=1)}),
        24,
        now=now,
    )[0] is False
    assert is_fresh(
        RawJobPosting(source="greenhouse", company_name="Example", discovered_at=now),
        24,
        now=now,
    ) == (False, None)
    body = cxs_request_body(limit=20, offset=0, search_text="software engineer")
    assert body["appliedFacets"] == {}
    assert list(body) == ["appliedFacets", "limit", "offset", "searchText"]


@pytest.mark.asyncio
async def test_fresh_job_protection_matrix(tmp_config) -> None:
    cases = (
        ("swe", "Software Engineer", "Required Qualifications\n0-2 years.\n" + WORK, "keep", None),
        ("acct", "Account Manager", "backend Python APIs production systems", "ROLE_MISMATCH", None),
        ("ai-thin", "AI Engineer", "Python, SQL, and APIs.", "SEMANTIC_REVIEW_UNAVAILABLE", "AI_ENGINEER"),
        ("ai-swe", "AI Engineer", "Required Qualifications\n0-2 years.\n" + WORK, "keep", None),
        ("qa", "QA Engineer", "Manual testing of business workflows.", "SEMANTIC_REVIEW_UNAVAILABLE", "QA_AUTOMATION_ENGINEER"),
    )
    for job_id, title, description, expected, family in cases:
        job = _job(
            job_id=job_id,
            job_title=title,
            description=description,
            direct_application_url=f"https://boards.greenhouse.io/acmerobotics/jobs/{job_id}",
        )
        assert is_fresh(job, tmp_config.freshness_hours)[0] is True
        state = PipelineState(config=tmp_config, jobs=[job], resources={"llm": NullLLMProvider()})
        await run_role_classification(state)
        if expected == "keep":
            assert [item.job_id for item in state.jobs] == [job_id]
            assert state.summary.semantic_review_required == 0
            assert state.summary.fresh_semantic_review_required == 0
            assert is_fresh(state.jobs[0], tmp_config.freshness_hours)[0] is True
        else:
            assert state.jobs == []
            assert rejection_code(state.rejected[0]) == expected
            if expected == "SEMANTIC_REVIEW_UNAVAILABLE":
                assert state.summary.rejected_by_role == 0
                assert state.summary.fresh_semantic_review_required == 1
                assert state.summary.fresh_semantic_review_unavailable == 1
                assert state.summary.semantic_by_family[family]["unavailable"] == 1
            else:
                assert state.summary.deterministic_rejects == 1
                assert state.summary.fresh_semantic_review_required == 0


@pytest.mark.asyncio
async def test_confidence_threshold_stays_at_0_6(tmp_config) -> None:
    low = ScriptLLM(_settings(retry_attempts=0), [_role_json(accept=True, confidence=0.59)])
    uncertain = PipelineState(
        config=tmp_config,
        jobs=[_job(job_id="low", job_title="AI Engineer", description="Python, SQL, and APIs.",
                   direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/low")],
        resources={"llm": low},
    )
    await run_role_classification(uncertain)
    assert uncertain.jobs == []
    assert rejection_code(uncertain.rejected[0]) == "SEMANTIC_REVIEW_UNCERTAIN"
    assert uncertain.summary.fresh_semantic_review_uncertain == 1

    exact = ScriptLLM(_settings(retry_attempts=0), [_role_json(accept=True, confidence=0.60)])
    accepted = PipelineState(
        config=tmp_config,
        jobs=[_job(job_id="exact", job_title="AI Engineer", description="Python, SQL, and APIs.",
                   direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/exact")],
        resources={"llm": exact},
    )
    await run_role_classification(accepted)
    assert [job.job_id for job in accepted.jobs] == ["exact"]
    assert accepted.summary.semantic_review_accepted == 1
    assert accepted.summary.fresh_semantic_review_accepted == 1


@pytest.mark.asyncio
async def test_circuit_blocks_then_recovers_and_reviews_the_next_job(tmp_config) -> None:
    provider = ScriptLLM(
        _settings(circuit_breaker_failures=1, circuit_reset_seconds=999, retry_attempts=0),
        [ValueError("down")],
    )
    assert provider.circuit_state == "CLOSED"
    blocked = PipelineState(
        config=tmp_config,
        jobs=[
            _job(job_id="amb", job_title="AI Engineer", description="Python, SQL, and APIs.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/amb"),
            _job(job_id="swe", job_title="Software Engineer",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/swe"),
            _job(job_id="held", job_title="AI Engineer", description="Python, SQL, and APIs.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/held"),
        ],
        resources={"llm": provider},
    )
    await run_role_classification(blocked)
    assert provider.circuit_state == "OPEN"
    assert provider.generated == 1
    assert [job.job_id for job in blocked.jobs] == ["swe"]
    assert {rejection_code(item) for item in blocked.rejected} == {"SEMANTIC_REVIEW_UNAVAILABLE"}
    assert blocked.summary.semantic_reviews_blocked_by_circuit == 1
    assert blocked.summary.semantic_review_unavailable == 2
    assert blocked.summary.deterministic_accepts == 1

    provider.settings = provider.settings.model_copy(update={"circuit_reset_seconds": 0})
    provider._script.extend([_role_json(accept=True, confidence=0.91), _role_json(accept=True, confidence=0.92)])
    recovered = PipelineState(
        config=tmp_config,
        jobs=[
            _job(job_id="probe", job_title="AI Engineer", description="Python, SQL, and APIs.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/probe"),
            _job(job_id="next", job_title="ML Engineer", description="Python, SQL, and APIs.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/next"),
        ],
        resources={"llm": provider},
    )
    await run_role_classification(recovered)
    assert provider.stats.recovery_attempts == 1
    assert provider.circuit_state == "CLOSED"
    assert [job.job_id for job in recovered.jobs] == ["probe", "next"]
    assert recovered.summary.semantic_review_accepted == 2


def test_manifest_separates_fresh_semantic_holds_from_role_mismatch(tmp_config) -> None:
    state = PipelineState(config=tmp_config)
    state.summary.jobs_discovered = 14670
    state.summary.jobs_after_cross_source_dedup = 14556
    state.summary.fresh_authoritative_jobs = 3
    state.summary.jobs_accepted = 0
    state.summary.deterministic_rejects = 3
    state.summary.semantic_review_required = 1700
    state.summary.semantic_review_unavailable = 1700
    state.summary.fresh_semantic_review_required = 0
    state.summary.fresh_semantic_review_unavailable = 0
    state.summary.llm_state_before = "OPEN"
    state.summary.llm_state_after = "OPEN"
    state.summary.exit_code = 0
    state.summary.xlsx_status = "written"
    state.summary.archive_status = "written"
    state.jobs = [_job(description="SECRET_DESCRIPTION_TEXT")]
    report = render_pipeline_health(state)
    assert "semantic_review_unavailable=1700" in report
    assert "fresh_semantic_review_required=0" in report
    assert "is a hold for missing model review and is not a role mismatch" in report
    assert "SECRET_DESCRIPTION_TEXT" not in report
    assert '"exit_code":0' in report
    assert '"source_empty":' in report


def test_empty_source_is_not_a_failure_and_a_cap_stays_partial() -> None:
    empty = SourceHealth(source="company_career", attempted=1, successful=1, failed=0, jobs_discovered=0)
    assert production_source_state(empty) == "EMPTY"
    assert source_completeness("company_career", empty) == "EMPTY"
    failed = SourceHealth(source="greenhouse", attempted=1, successful=0, failed=1, jobs_discovered=0)
    assert production_source_state(failed) == "FAILED"
    capped = SourceHealth(source="workday", attempted=2, successful=2, failed=0, jobs_discovered=2000)
    assert source_completeness("workday", capped, {"by_source": {"workday": {"caps": 2}}}) == "PARTIAL"


@pytest.mark.asyncio
async def test_one_source_failure_does_not_stop_the_others(tmp_config) -> None:
    async def greenhouse():
        raise RuntimeError("board down")

    async def workday():
        return [_posting("workday", "w1", method="cxs", title="Software Engineer")]

    result = await orchestrate(
        tmp_config,
        collectors={"greenhouse": greenhouse, "workday": workday},
        llm=NullLLMProvider(),
        profile=_profile(),
    )
    assert result.sources["greenhouse"].status == "FAILED"
    assert result.sources["workday"].status == "SUCCESS"
    assert [job.job_id for job in result.funnel.state.jobs] == ["w1"]


@pytest.mark.asyncio
async def test_nonzero_fixture_writes_a_schema_stable_row(tmp_config) -> None:
    job = _job()
    state = PipelineState(config=tmp_config, jobs=[job], resources={"llm": NullLLMProvider()})
    await run_role_classification(state)
    assert state.summary.deterministic_accepts == 1
    tmp_config.dry_run = False
    tmp_config.send_email = False
    await run_output(state)
    assert state.summary.xlsx_status == "written"
    assert state.summary.archive_status == "written"
    current = tmp_config.path(tmp_config.settings.output.current_workbook)
    workbook = load_workbook(current)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(ALL_COLUMNS)
    assert sheet.max_row == 2
    row = {headers[i]: sheet.cell(2, i + 1).value for i in range(len(headers))}
    assert row["Job ID"] == "swe"
    assert row["Direct Application URL"] == job.direct_application_url
    assert "Applied" in headers and "Status" in headers
    assert not any("semantic" in str(header).lower() or str(header).startswith("LLM") for header in headers)
    workbook.close()


def test_zero_candidate_run_does_not_copy_previous_jobs(tmp_path) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks([make_job(job_id="yesterday")], current_path=current, archive_dir=archive, write_archive=False)
    existing = list(iter_workbook_rows(current))
    write_workbooks([], current_path=current, archive_dir=archive, existing_rows=existing, write_archive=False)
    workbook = load_workbook(current)
    assert workbook.active.max_row == 1
    workbook.close()
    assert list(iter_workbook_rows(current)) == []


def test_archive_failure_preserves_current_and_previous_archive(tmp_path, monkeypatch) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks([make_job(job_id="kept")], current_path=current, archive_dir=archive, write_archive=True)
    current_bytes = current.read_bytes()
    previous = next(archive.glob("*.xlsx"))
    previous_bytes = previous.read_bytes()
    real_replace = os.replace

    def fail_archive(src, dst):
        if Path(dst) == previous:
            raise OSError("archive disk full")
        return real_replace(src, dst)

    monkeypatch.setattr("src.services.xlsx.os.replace", fail_archive)
    with pytest.raises(OSError, match="archive disk full"):
        write_workbooks([make_job(job_id="new")], current_path=current, archive_dir=archive, write_archive=True)
    assert current.read_bytes() == current_bytes
    assert previous.read_bytes() == previous_bytes
    assert hashlib.sha256(previous.read_bytes()).hexdigest() == hashlib.sha256(previous_bytes).hexdigest()


@pytest.mark.asyncio
async def test_archive_failure_does_not_report_the_workbook_written(tmp_config, monkeypatch) -> None:
    tmp_config.dry_run = False
    tmp_config.send_email = False
    current = tmp_config.path(tmp_config.settings.output.current_workbook)
    archive = tmp_config.path(tmp_config.settings.output.archive_dir)
    write_workbooks([make_job(job_id="kept")], current_path=current, archive_dir=archive, write_archive=True)
    before = current.read_bytes()

    def fail_archive(src, dst):
        if Path(dst).parent == archive and not str(dst).endswith(".tmp"):
            raise OSError("archive disk full")
        return os.replace(src, dst)

    monkeypatch.setattr("src.services.xlsx.os.replace", fail_archive)
    state = PipelineState(config=tmp_config, jobs=[make_job(job_id="new")])
    with pytest.raises(OSError):
        await run_output(state)
    assert current.read_bytes() == before
    assert state.summary.xlsx_status == "not_written"
    assert state.summary.archive_status == "not_written"
    assert state.summary.email_status == "not attempted"


def test_security_and_workflow_boundaries_remain() -> None:
    workflow = (ROOT / ".github" / "workflows" / "daily_jobs.yml").read_text(encoding="utf-8")
    assert 'cron: "0 12 * * *"' in workflow
    assert 'python-version: "3.12"' in workflow
    assert "python -m src.main" in workflow
    assert "contents: write" in workflow
    assert "git add data/current data/archive" in workflow
    assert "resume.pdf" not in workflow
    assert "--multi-source" not in workflow
    assert "data/candidate/" in (ROOT / ".gitignore").read_text(encoding="utf-8")
    settings = yaml.safe_load((ROOT / "config" / "settings.yaml").read_text(encoding="utf-8"))
    assert settings["discovery"]["sources"]["jobright"]["enabled"] is False
    assert settings["discovery"]["sources"]["workday"]["browser_enabled"] is False
    assert settings["discovery"]["persist_ats_registry"] is False
