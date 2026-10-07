"""Phase 16: public repository and GitHub Actions boundaries."""

from __future__ import annotations

import hashlib
from pathlib import Path

import yaml
from openpyxl import load_workbook

from src.agents.intelligence_agent import run_job_intelligence
from src.main import build_parser
from src.models.config import load_config
from src.models.state import PipelineState
from src.services.resume_store import load_profile
from src.services.xlsx import ALL_COLUMNS

ROOT = Path(__file__).resolve().parents[1]
HISTORICAL_ARCHIVES = {
    # Bytes already committed by the hosted daily workflow on origin/main.
    "2026-09-24.xlsx": "547f322803b3ba97adda050f77144b4010ce43c35038f966d439c57df851e5b4",
    "2026-09-25.xlsx": "3bde8122e166a7f3fc8c89217619f4f8f5eae4fb462eaecb38fcc3ddaf094747",
    "2026-09-26.xlsx": "cc3c7067f528b393182654ccba7e157f7df30a2927fb40d950d30fae7a5bda8a",
}
FORBIDDEN_HOSTS = ("linkedin.com", "indeed.com", "glassdoor.com", "jobright.ai", "ziprecruiter.com")


def test_workflow_is_production_only_and_least_privilege() -> None:
    workflow = (ROOT / ".github" / "workflows" / "daily_jobs.yml").read_text(encoding="utf-8")
    assert 'cron: "0 12 * * *"' in workflow
    assert 'python-version: "3.12"' in workflow
    assert "python -m src.main" in workflow
    assert "permissions:\n  contents: write\n" in workflow
    for forbidden in (
        "pull-requests:",
        "issues:",
        "packages:",
        "deployments:",
        "id-token:",
        "actions: write",
    ):
        assert forbidden not in workflow
    run = workflow.split("name: Run pipeline", 1)[1].split("name: Commit generated tracker", 1)[0]
    # Command may pipe through tee for log capture; core invocation must not carry flags.
    assert "python -m src.main" in run
    assert "python -m src.main --" not in run
    assert "resume.pdf" not in workflow
    assert "data/candidate" not in workflow
    assert "git add data/current data/archive" in workflow
    assert "ANTHROPIC_API_KEY: ${{ secrets.ANTHROPIC_API_KEY }}" in workflow
    assert "ANTHROPIC_API_KEY=" not in workflow
    assert "GEMINI_API_KEY: ${{ secrets.GEMINI_API_KEY }}" in workflow
    assert "GEMINI_API_KEY=" not in workflow


def test_gitignore_protects_private_files_and_keeps_the_project() -> None:
    ignored = (ROOT / ".gitignore").read_text(encoding="utf-8")
    for required in (
        ".venv/",
        ".env",
        "!.env.example",
        "__pycache__/",
        ".pytest_cache/",
        "*.log",
        "data/cache/",
        "data/navigation/",
        "data/reports/",
        "data/candidate/",
        "~$*.xlsx",
        "*.xlsx.tmp",
    ):
        assert required in ignored
    for kept in ("src/", "tests/", "config/", ".github/"):
        assert kept not in ignored.split()


def test_production_configuration_is_safe_and_deterministic() -> None:
    settings_path = ROOT / "config" / "settings.yaml"
    before = {
        path.name: path.read_bytes()
        for path in (ROOT / "config").glob("*.yaml")
    }
    config = load_config(ROOT / "config", send_email=False)
    assert config.settings.discovery.sources.jobright.enabled is False
    assert config.settings.discovery.sources.workday.browser_enabled is False
    assert config.settings.discovery.persist_ats_registry is False
    assert config.freshness_hours == 24
    assert config.fixture_mode is False
    for path in (ROOT / "config").glob("*.yaml"):
        text = path.read_text(encoding="utf-8")
        assert "AIza" not in text
        assert "smtp_password:" not in text.lower()
        assert "C:\\" not in text
        assert "/Users/" not in text
        assert path.read_bytes() == before[path.name]
    example = (ROOT / ".env.example").read_text(encoding="utf-8")
    assert "ANTHROPIC_API_KEY=\n" in example
    assert "GEMINI_API_KEY=\n" in example
    assert "SMTP_PASSWORD=\n" in example
    parsed = yaml.safe_load(settings_path.read_text(encoding="utf-8"))
    assert parsed["discovery"]["sources"]["jobright"]["enabled"] is False


def test_default_cli_does_not_inherit_smoke_behavior() -> None:
    settings = (ROOT / "config" / "settings.yaml").read_bytes()
    default = build_parser().parse_args([])
    assert default.fixture_mode is False
    assert default.dry_run is False
    assert default.multi_source_smoke_test is False
    assert default.multi_source_browser_smoke_test is False
    assert default.jobright_smoke_test is False
    assert default.workday_browser_smoke_test is False
    smoke = build_parser().parse_args(
        ["--multi-source-browser-smoke-test", "--sources", "greenhouse,workday", "--dry-run", "--no-email"]
    )
    assert smoke.multi_source_browser_smoke_test is True
    assert default.multi_source_browser_smoke_test is False
    assert (ROOT / "config" / "settings.yaml").read_bytes() == settings


def test_missing_resume_does_not_block_the_daily_pipeline(tmp_path, monkeypatch) -> None:
    missing = tmp_path / "missing-profile.json"
    assert load_profile(missing) is None
    config = load_config(ROOT / "config", dry_run=True, send_email=False)
    candidate = config.settings.candidate.model_copy(
        update={"resume_path": "missing-resume.pdf", "profile_path": str(missing)}
    )
    config.settings = config.settings.model_copy(update={"candidate": candidate})
    seen: list[Path] = []
    real_load = load_profile

    def _spy(path: Path):
        seen.append(Path(path))
        return real_load(path)

    monkeypatch.setattr("src.services.resume_store.load_profile", _spy)
    state = PipelineState(config=config, jobs=[])
    import asyncio

    asyncio.run(run_job_intelligence(state))
    assert seen == [missing]
    assert state.summary.intelligence_evaluated == 0
    assert state.summary.critic_reviewed == 0


def test_public_workbooks_keep_the_schema_and_historical_bytes() -> None:
    current = ROOT / "data" / "current" / "jobs.xlsx"
    _assert_public_workbook(current)
    archive = ROOT / "data" / "archive"
    for name, digest in HISTORICAL_ARCHIVES.items():
        path = archive / name
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
        _assert_public_workbook(path)
    today = archive / "2026-09-27.xlsx"
    if today.exists():
        _assert_public_workbook(today)


def test_identical_workbook_writes_stay_one_row(tmp_path) -> None:
    from tests.conftest import make_job
    from src.services.xlsx import iter_workbook_rows, write_workbooks

    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    job = make_job(job_id="same", job_title="Software Engineer")
    write_workbooks([job], current_path=current, archive_dir=archive, write_archive=True)
    rows = list(iter_workbook_rows(current))
    rows[0]["Applied"] = "☑ Applied"
    rows[0]["Status"] = "Interview"
    write_workbooks([job], current_path=current, archive_dir=archive, existing_rows=rows, write_archive=False)
    again = list(iter_workbook_rows(current))
    assert len(again) == 1
    assert again[0]["Job ID"] == "same"
    assert again[0]["Applied"] == "☑ Applied"
    assert again[0]["Status"] == "Interview"


def _assert_public_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers == list(ALL_COLUMNS)
    assert "Applied" in headers and "Status" in headers
    joined = " ".join(str(header) for header in headers).lower()
    assert "description" not in joined
    assert "semantic" not in joined
    assert "llm" not in joined
    url_index = headers.index("Direct Application URL")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        url = str(row[url_index] or "").lower()
        assert not any(host in url for host in FORBIDDEN_HOSTS)
        blob = " ".join(str(value or "") for value in row).lower()
        assert "aiza" not in blob
        assert "smtp" not in blob
    workbook.close()
