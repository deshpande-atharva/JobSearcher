"""Phase 14: semantic-review observability, email, archives, and exit codes."""

from __future__ import annotations

import hashlib
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
import yaml

from src.agents.dedup_agent import run_dedup
from src.agents.freshness_agent import run_freshness
from src.agents.h1b_sponsorship_agent import run_h1b_enrichment
from src.agents.location_agent import run_location_employment
from src.agents.output_agent import run_output
from src.agents.qc_agent import run_quality_control
from src.agents.role_agent import run_role_classification
from src.agents.seniority_agent import run_seniority
from src.agents.url_agent import run_url_verification
from src.graph.pipeline import build_graph
from src.llm.base import NullLLMProvider
from src.llm.schemas import RoleClassificationResult
from src.main import async_main
from src.models.config import Secrets
from src.models.job import DateSource, EmploymentType, RawJobPosting, RejectionReason
from src.models.state import PipelineState, SourceHealth
from src.services.fresh_preview import attach_fresh_preview
from src.services.freshness import is_fresh
from src.services.notifications import build_email_body, build_email_subject, send_run_summary
from src.services.production_report import production_source_state, render_pipeline_health
from src.services.rejection_codes import rejection_code
from src.services.roles import classify_role
from src.services.xlsx import ALL_COLUMNS, iter_workbook_rows, write_workbooks
from src.sources.workday import cxs_request_body
from src.utils.dates import utcnow
from tests.conftest import make_job
from tests.test_llm_reliability import ProviderError, ScriptLLM, _settings

WORK = "Write production code and run code review for backend services."
ROOT = Path(__file__).resolve().parents[1]


def _disable_visa(config) -> None:
    visa = config.settings.visa.model_copy(update={"enabled": False})
    config.settings = config.settings.model_copy(update={"visa": visa})


def _role_json(*, accept: bool, confidence: float, family: str = "MACHINE_LEARNING_ENGINEER") -> str:
    import json

    return json.dumps(
        {
            "is_software_engineering": accept,
            "role_family": family if accept else "NOT_SOFTWARE",
            "confidence": confidence,
            "reasoning": "fixture decision",
        }
    )


async def _qualify(config, job, llm=None):
    _disable_visa(config)
    state = PipelineState(config=config, jobs=[job], resources={"llm": llm or NullLLMProvider()})
    await run_role_classification(state)
    await run_seniority(state)
    await run_location_employment(state)
    await run_freshness(state)
    await run_url_verification(state)
    await run_h1b_enrichment(state)
    await run_dedup(state)
    await run_quality_control(state)
    return state


def _swe(**overrides):
    data = dict(
        job_title="Software Engineer",
        location="Boston, MA",
        description="Required Qualifications\n0-2 years of experience.\n" + WORK,
        posted_at=utcnow() - timedelta(hours=2),
        date_source=DateSource.POSTED_DATE,
        employment_type=EmploymentType.FULL_TIME,
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/swe",
        job_id="swe",
    )
    data.update(overrides)
    return make_job(**data)


def test_graph_order_is_unchanged() -> None:
    edges = {(edge.source, edge.target) for edge in build_graph().get_graph().edges}
    assert ("seniority", "location") in edges
    assert ("location", "freshness") in edges
    assert ("freshness", "url") in edges


def test_classification_states_match_the_existing_taxonomy(tmp_config) -> None:
    roles = tmp_config.roles
    swe = classify_role("Software Engineer", "", roles)
    assert swe.classification_state == "DETERMINISTIC_ACCEPT"
    assert swe.needs_llm is False
    account = classify_role(
        "Account Manager",
        "backend Python APIs production systems",
        roles,
    )
    assert account.classification_state == "DETERMINISTIC_REJECT"
    assert account.needs_llm is False
    scientist = classify_role("Data Scientist", WORK, roles)
    assert scientist.classification_state == "DETERMINISTIC_REJECT"
    assert scientist.needs_llm is False
    ambiguous = classify_role("AI Engineer", "Python, SQL, and APIs.", roles)
    assert ambiguous.classification_state == "SEMANTIC_REVIEW_REQUIRED"
    assert ambiguous.needs_llm is True
    promoted = classify_role("AI Engineer", WORK, roles)
    assert promoted.classification_state == "DETERMINISTIC_ACCEPT"
    assert promoted.needs_llm is False
    solutions = classify_role("Solutions Engineer", "Pre-sales demos and proposals.", roles)
    assert solutions.classification_state == "DETERMINISTIC_ACCEPT"
    assert solutions.family == "PRODUCT_ENGINEER"
    for title in ("DevOps Engineer", "Site Reliability Engineer", "SDET", "Automation Engineer"):
        verdict = classify_role(title, "", roles)
        assert verdict.classification_state == "DETERMINISTIC_ACCEPT", title
        assert verdict.needs_llm is False
    for title in ("QA Engineer", "Test Engineer"):
        manual = classify_role(title, "Manual testing of business workflows.", roles)
        assert manual.classification_state == "SEMANTIC_REVIEW_REQUIRED", title
        clear = classify_role(title, WORK, roles)
        assert clear.classification_state == "DETERMINISTIC_ACCEPT", title


def test_freshness_boundary_is_unchanged() -> None:
    now = datetime(2026, 9, 27, 12, tzinfo=timezone.utc)
    posted = RawJobPosting(
        source="greenhouse",
        company_name="Example",
        title="Software Engineer",
        posted_at=now - timedelta(hours=24),
        date_source=DateSource.POSTED_DATE,
        discovered_at=now,
    )
    assert is_fresh(posted, 24, now=now)[0] is True
    assert is_fresh(posted.model_copy(update={"posted_at": now - timedelta(hours=24, seconds=1)}), 24, now=now)[0] is False
    assert is_fresh(RawJobPosting(source="greenhouse", company_name="Example", discovered_at=now), 24, now=now) == (False, None)


def test_workday_public_body_is_unchanged() -> None:
    body = cxs_request_body(limit=20, offset=0, search_text="")
    assert list(body) == ["appliedFacets", "limit", "offset", "searchText"]
    assert body["appliedFacets"] == {}
    assert body["limit"] == 20


@pytest.mark.asyncio
async def test_deterministic_jobs_ignore_an_unavailable_model(tmp_config) -> None:
    kept = await _qualify(tmp_config, _swe())
    assert [job.job_id for job in kept.jobs] == ["swe"]
    assert kept.summary.semantic_review_required == 0
    rejected = await _qualify(
        tmp_config,
        _swe(
            job_id="acct",
            job_title="Account Manager",
            description="backend Python APIs production systems",
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/acct",
        ),
    )
    assert rejected.jobs == []
    assert rejection_code(rejected.rejected[0]) == "ROLE_MISMATCH"
    assert rejected.summary.semantic_review_unavailable == 0


@pytest.mark.asyncio
async def test_ambiguous_ai_without_gemini_is_not_a_role_mismatch(tmp_config) -> None:
    state = await _qualify(
        tmp_config,
        _swe(
            job_id="ai-open",
            job_title="AI Engineer",
            description="Required Qualifications\n0-2 years of experience.\nPython, SQL, and APIs.",
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-open",
        ),
    )
    assert state.jobs == []
    assert rejection_code(state.rejected[0]) == "SEMANTIC_REVIEW_UNAVAILABLE"
    assert state.summary.rejected_by_role == 0
    assert state.summary.semantic_review_required == 1
    assert state.summary.semantic_review_unavailable == 1
    assert state.summary.semantic_by_family["AI_ENGINEER"]["unavailable"] == 1


@pytest.mark.asyncio
async def test_ai_with_software_signals_qualifies_without_gemini(tmp_config) -> None:
    state = await _qualify(
        tmp_config,
        _swe(
            job_id="ai-swe",
            job_title="AI Engineer",
            description="Required Qualifications\n0-2 years of experience.\n" + WORK,
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-swe",
        ),
    )
    assert [job.job_id for job in state.jobs] == ["ai-swe"]
    assert state.summary.semantic_review_required == 0


@pytest.mark.asyncio
async def test_semantic_review_accept_reject_and_uncertain(tmp_config) -> None:
    accept = ScriptLLM(_settings(retry_attempts=0), [_role_json(accept=True, confidence=0.91)])
    accepted = await _qualify(
        tmp_config,
        _swe(
            job_id="ai-yes",
            job_title="AI Engineer",
            description="Python, SQL, and APIs. Required Qualifications\n0-2 years of experience.",
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-yes",
        ),
        llm=accept,
    )
    assert [job.job_id for job in accepted.jobs] == ["ai-yes"]
    assert accepted.summary.semantic_review_accepted == 1
    assert accepted.jobs[0].direct_application_url.startswith("https://boards.greenhouse.io/")

    reject = ScriptLLM(_settings(retry_attempts=0), [_role_json(accept=False, confidence=0.93)])
    rejected = await _qualify(
        tmp_config,
        _swe(
            job_id="ai-no",
            job_title="AI Engineer",
            description="Python, SQL, and APIs.",
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-no",
        ),
        llm=reject,
    )
    assert rejected.jobs == []
    assert rejection_code(rejected.rejected[0]) == "ROLE_MISMATCH"
    assert rejected.summary.semantic_review_rejected == 1

    unsure = ScriptLLM(_settings(retry_attempts=0), [_role_json(accept=False, confidence=0.2)])
    uncertain = await _qualify(
        tmp_config,
        _swe(
            job_id="ai-maybe",
            job_title="AI Engineer",
            description="Python, SQL, and APIs.",
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-maybe",
        ),
        llm=unsure,
    )
    assert uncertain.jobs == []
    assert rejection_code(uncertain.rejected[0]) == "SEMANTIC_REVIEW_UNCERTAIN"
    assert uncertain.summary.rejected_by_role == 0
    assert uncertain.summary.semantic_review_uncertain == 1


@pytest.mark.asyncio
async def test_llm_cannot_rewrite_posting_identity(tmp_config, monkeypatch) -> None:
    job = _swe(
        job_id="ai-id",
        job_title="AI Engineer",
        description="Python, SQL, and APIs.",
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-id",
    )
    original = (job.job_id, job.source, job.direct_application_url, job.posted_at, job.date_source, job.location)

    async def evil(current, llm, state):
        current.job_id = "changed"
        current.source = "jobright"
        current.direct_application_url = "https://www.linkedin.com/jobs/1"
        current.posted_at = None
        current.date_source = DateSource.DISCOVERED_DATE
        current.location = "London, United Kingdom"
        return RoleClassificationResult(
            is_software_engineering=True,
            role_family="SOFTWARE_ENGINEER",
            confidence=0.95,
            reasoning="override",
        )

    monkeypatch.setattr("src.agents.role_agent._ask_llm", evil)
    provider = ScriptLLM(_settings(retry_attempts=0), [])
    state = PipelineState(config=tmp_config, jobs=[job], resources={"llm": provider})
    await run_role_classification(state)
    assert [item.job_id for item in state.jobs] == ["ai-id"]
    kept = state.jobs[0]
    assert (kept.job_id, kept.source, kept.direct_application_url, kept.posted_at, kept.date_source, kept.location) == original


@pytest.mark.asyncio
async def test_provider_failures_are_semantic_review_unavailable(tmp_config) -> None:
    cases = (
        ("429", [ProviderError(429, "429 RESOURCE_EXHAUSTED"), ProviderError(429, "429 RESOURCE_EXHAUSTED")], {}),
        ("503", [ProviderError(503, "503 UNAVAILABLE")], {"circuit_breaker_failures": 1, "retry_attempts": 0}),
        ("timeout", [TimeoutError("timed out")], {"retry_attempts": 0}),
        ("schema", ["not-json"], {"retry_attempts": 0}),
    )
    for name, script, options in cases:
        provider = ScriptLLM(_settings(**options), script)
        state = PipelineState(
            config=tmp_config,
            jobs=[
                _swe(
                    job_id=name,
                    job_title="AI Engineer",
                    description="Python, SQL, and APIs.",
                    direct_application_url=f"https://boards.greenhouse.io/acmerobotics/jobs/{name}",
                )
            ],
            resources={"llm": provider},
        )
        await run_role_classification(state)
        assert state.jobs == [], name
        assert rejection_code(state.rejected[0]) == "SEMANTIC_REVIEW_UNAVAILABLE", name
        assert state.summary.rejected_by_role == 0, name


@pytest.mark.asyncio
async def test_open_circuit_blocks_semantic_review_and_a_probe_can_accept(tmp_config) -> None:
    blocked = ScriptLLM(
        _settings(circuit_breaker_failures=1, circuit_reset_seconds=999, retry_attempts=0),
        [ValueError("down"), _role_json(accept=True, confidence=0.9)],
    )
    state = PipelineState(
        config=tmp_config,
        jobs=[
            _swe(job_id="first", job_title="AI Engineer", description="Python, SQL, and APIs.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/first"),
            _swe(job_id="second", job_title="ML Engineer", description="Python, SQL, and APIs.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/second"),
        ],
        resources={"llm": blocked},
    )
    await run_role_classification(state)
    assert blocked.circuit_state == "OPEN"
    assert blocked.generated == 1
    assert state.summary.semantic_reviews_blocked_by_circuit == 1
    assert state.summary.semantic_review_unavailable == 2
    assert {rejection_code(item) for item in state.rejected} == {"SEMANTIC_REVIEW_UNAVAILABLE"}

    recovering = ScriptLLM(
        _settings(circuit_breaker_failures=1, circuit_reset_seconds=0, retry_attempts=0),
        [ValueError("down"), _role_json(accept=True, confidence=0.9)],
    )
    recovered = PipelineState(
        config=tmp_config,
        jobs=[
            _swe(job_id="probe-a", job_title="AI Engineer", description="Python, SQL, and APIs. Required Qualifications\n0-2 years.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/probe-a"),
            _swe(job_id="probe-b", job_title="AI Engineer", description="Python, SQL, and APIs. Required Qualifications\n0-2 years.",
                 direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/probe-b"),
        ],
        resources={"llm": recovering},
    )
    await run_role_classification(recovered)
    assert recovering.circuit_state == "CLOSED"
    assert [job.job_id for job in recovered.jobs] == ["probe-b"]
    assert recovered.summary.semantic_review_accepted == 1


@pytest.mark.asyncio
async def test_qa_semantic_fallback_with_gemini_unavailable(tmp_config) -> None:
    manual = await _qualify(
        tmp_config,
        _swe(job_id="qa", job_title="QA Engineer", description="Manual testing of business workflows.",
             direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/qa"),
    )
    assert rejection_code(manual.rejected[0]) == "SEMANTIC_REVIEW_UNAVAILABLE"
    assert manual.summary.semantic_by_family["QA_AUTOMATION_ENGINEER"]["unavailable"] == 1
    clear = await _qualify(
        tmp_config,
        _swe(
            job_id="qa-clear",
            job_title="Test Engineer",
            description="Required Qualifications\n0-2 years of experience.\n" + WORK,
            direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/qa-clear",
        ),
    )
    assert [job.job_id for job in clear.jobs] == ["qa-clear"]
    assert clear.summary.semantic_review_required == 0


@pytest.mark.asyncio
async def test_semantic_fresh_audit_omits_the_description(tmp_config) -> None:
    job = _swe(
        job_id="ai-fresh",
        job_title="AI Engineer",
        location="Boston, MA",
        description="SECRET_DESCRIPTION_TEXT Python, SQL, and APIs.",
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/ai-fresh",
    )
    state = PipelineState(
        config=tmp_config,
        jobs=[job],
        raw_postings=[
            RawJobPosting(
                source="greenhouse",
                company_name=job.company,
                title=job.job_title,
                job_id=job.job_id,
                apply_url=job.direct_application_url,
                location_raw=job.location,
                description=job.description,
                posted_at=job.posted_at,
                date_source=DateSource.POSTED_DATE,
            )
        ],
        resources={"llm": NullLLMProvider()},
    )
    await run_role_classification(state)
    attach_fresh_preview(state)
    report = render_pipeline_health(state)
    assert "semantic_fresh_job:" in report
    assert "semantic_family=AI_ENGINEER" in report
    assert "semantic_status=SEMANTIC_REVIEW_UNAVAILABLE" in report
    assert "SECRET_DESCRIPTION_TEXT" not in report
    assert "run_manifest:" in report
    assert "password" not in report


def test_source_states_cover_partial_and_failed() -> None:
    partial = SourceHealth(source="workday", attempted=2, successful=1, failed=1, jobs_discovered=10)
    failed = SourceHealth(source="workday", attempted=2, successful=0, failed=2, jobs_discovered=0)
    assert production_source_state(partial) == "PARTIAL"
    assert production_source_state(failed) == "FAILED"


def test_email_subject_and_body_omit_sensitive_content() -> None:
    from src.models.state import RunSummary

    finished = utcnow()
    run = RunSummary(
        jobs_accepted=1,
        jobs_discovered=12,
        fresh_authoritative_jobs=3,
        llm_circuit_state="OPEN",
        llm_state_after="OPEN",
        run_started_at=finished - timedelta(seconds=30),
        run_finished_at=finished,
    )
    job = make_job(description="SECRET_DESCRIPTION_TEXT resume body")
    body = build_email_body(run, [job])
    subject = build_email_subject(run)
    assert "status=completed" in subject
    assert finished.strftime("%Y-%m-%d") in subject
    assert "Discovered jobs: 12" in body
    assert "Fresh jobs: 3" in body
    assert "Final candidates: 1" in body
    assert "LLM state: OPEN" in body
    assert "Runtime:" in body
    assert "SECRET_DESCRIPTION_TEXT" not in body
    assert "smtp_password" not in body
    zero = build_email_body(RunSummary(jobs_accepted=0), [])
    assert "valid zero-result run" in zero
    assert "Final jobs: 0" in zero


def test_smtp_missing_is_explicit_and_nonfatal(tmp_config) -> None:
    tmp_config.send_email = True
    tmp_config.dry_run = False
    status = send_run_summary(tmp_config, __import__("src.models.state", fromlist=["RunSummary"]).RunSummary())
    assert status == "skipped (SMTP_NOT_CONFIGURED)"


def test_smtp_success_and_failure_use_a_double(tmp_config, monkeypatch) -> None:
    from src.models.state import RunSummary

    tmp_config.send_email = True
    tmp_config.dry_run = False
    tmp_config.secrets = Secrets(
        smtp_host="smtp.example.com",
        smtp_port=587,
        smtp_username="user",
        smtp_password="smtp-secret-value",
        notification_email="jobs@example.com",
    )
    sent: list = []

    class FakeSMTP:
        def __init__(self, *args, **kwargs) -> None:
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args) -> None:
            return None

        def ehlo(self) -> None:
            return None

        def starttls(self) -> None:
            return None

        def login(self, username, password) -> None:
            return None

        def send_message(self, message) -> None:
            sent.append(message)

    monkeypatch.setattr("src.services.notifications.smtplib.SMTP", FakeSMTP)
    summary = RunSummary(jobs_accepted=1, jobs_discovered=4, fresh_authoritative_jobs=1, run_finished_at=utcnow())
    job = make_job(description="SECRET_DESCRIPTION_TEXT")
    assert send_run_summary(tmp_config, summary, [job]) == "sent"
    payload = sent[0].get_content()
    assert "smtp-secret-value" not in payload
    assert "SECRET_DESCRIPTION_TEXT" not in payload
    assert "status=completed" in sent[0]["Subject"]

    class BrokenSMTP(FakeSMTP):
        def __enter__(self):
            raise OSError("connection refused")

    monkeypatch.setattr("src.services.notifications.smtplib.SMTP", BrokenSMTP)
    assert send_run_summary(tmp_config, summary, [job]).startswith("failed")


@pytest.mark.asyncio
async def test_email_failure_does_not_fail_the_pipeline(tmp_config, monkeypatch) -> None:
    tmp_config.dry_run = True
    tmp_config.send_email = True
    tmp_config.secrets = Secrets(
        smtp_host="smtp.example.com",
        smtp_port=587,
        smtp_username="user",
        smtp_password="smtp-secret-value",
        notification_email="jobs@example.com",
    )

    class BrokenSMTP:
        def __init__(self, *args, **kwargs) -> None:
            pass

        def __enter__(self):
            raise OSError("connection refused")

        def __exit__(self, *args) -> None:
            return None

    monkeypatch.setattr("src.services.notifications.smtplib.SMTP", BrokenSMTP)
    state = PipelineState(config=tmp_config, jobs=[])
    await run_output(state)
    assert state.summary.email_failed is True
    assert state.summary.email_status.startswith("failed")


@pytest.mark.asyncio
async def test_xlsx_failure_propagates_and_pipeline_exit_is_nonzero(tmp_config, monkeypatch) -> None:
    tmp_config.dry_run = False
    tmp_config.send_email = False

    def boom(*args, **kwargs):
        raise OSError("disk full")

    monkeypatch.setattr("src.agents.output_agent.write_workbooks", boom)
    with pytest.raises(OSError):
        await run_output(PipelineState(config=tmp_config, jobs=[]))

    async def crash(config):
        raise OSError("disk full")

    monkeypatch.setattr("src.graph.pipeline.run_pipeline", crash)
    assert await async_main(["--dry-run", "--no-email"]) == 1


@pytest.mark.asyncio
async def test_valid_zero_run_exits_zero_and_bad_config_exits_nonzero(monkeypatch) -> None:
    async def empty(config):
        state = PipelineState(config=config)
        state.summary.jobs_accepted = 0
        state.summary.run_finished_at = utcnow()
        return state

    monkeypatch.setattr("src.graph.pipeline.run_pipeline", empty)
    assert await async_main(["--dry-run", "--no-email"]) == 0
    assert await async_main(["--config-dir", "config-does-not-exist", "--no-email"]) == 2


def test_workbook_schema_and_daily_identity(tmp_path) -> None:
    assert len(ALL_COLUMNS) == 17
    assert "Applied" in ALL_COLUMNS and "Status" in ALL_COLUMNS
    assert not any("semantic" in column.lower() or column.startswith("LLM") for column in ALL_COLUMNS)
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    original = make_job(job_id="same", job_title="Software Engineer")
    write_workbooks([original], current_path=current, archive_dir=archive, write_archive=False)
    from openpyxl import load_workbook
    from src.models.job import ApplicationStatus, AppliedFlag

    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("Applied") + 1, AppliedFlag.APPLIED.value)
    sheet.cell(2, headers.index("Status") + 1, ApplicationStatus.INTERVIEW.value)
    wb.save(current)
    wb.close()
    existing = list(iter_workbook_rows(current))
    updated = make_job(job_id="same", job_title="Software Engineer II")
    extra = make_job(
        job_id="other",
        job_title="Software Engineer II",
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/other",
    )
    write_workbooks([updated, extra], current_path=current, archive_dir=archive, existing_rows=existing, write_archive=False)
    rows = list(iter_workbook_rows(current))
    assert len(rows) == 2
    by_id = {row["Job ID"]: row for row in rows}
    assert by_id["same"]["Job Title"] == "Software Engineer II"
    assert by_id["same"]["Applied"] == AppliedFlag.APPLIED.value
    assert by_id["same"]["Status"] == ApplicationStatus.INTERVIEW.value
    assert by_id["other"]["Job Title"] == "Software Engineer II"


def test_previous_archive_is_unchanged_when_a_new_day_is_written(tmp_path, monkeypatch) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    clock = {"now": datetime(2026, 9, 27, tzinfo=timezone.utc)}
    monkeypatch.setattr("src.services.xlsx.utcnow", lambda: clock["now"])
    write_workbooks([make_job(job_id="day27")], current_path=current, archive_dir=archive, write_archive=True)
    previous = archive / "2026-09-27.xlsx"
    digest = hashlib.sha256(previous.read_bytes()).hexdigest()
    clock["now"] = datetime(2026, 9, 28, tzinfo=timezone.utc)
    write_workbooks([make_job(job_id="day28")], current_path=current, archive_dir=archive, write_archive=True)
    assert hashlib.sha256(previous.read_bytes()).hexdigest() == digest
    assert (archive / "2026-09-28.xlsx").is_file()


def test_github_workflow_and_privacy_boundaries() -> None:
    workflow = (ROOT / ".github" / "workflows" / "daily_jobs.yml").read_text(encoding="utf-8")
    assert 'cron: "0 12 * * *"' in workflow
    assert 'python-version: "3.12"' in workflow
    assert "python -m src.main" in workflow
    assert "git add data/current data/archive" in workflow
    assert "contents: write" in workflow
    assert "--multi-source" not in workflow
    assert "resume.pdf" not in workflow
    assert "data/candidate" not in workflow
    gitignore = (ROOT / ".gitignore").read_text(encoding="utf-8")
    assert "data/candidate/" in gitignore
    settings = yaml.safe_load((ROOT / "config" / "settings.yaml").read_text(encoding="utf-8"))
    assert settings["discovery"]["sources"]["jobright"]["enabled"] is False
    assert settings["discovery"]["sources"]["workday"]["browser_enabled"] is False
    assert settings["discovery"]["persist_ats_registry"] is False


class _Probe:
    def __init__(self, ok: bool, status: int = 200) -> None:
        self.ok = ok
        self.status = status
        self.error = None if ok else "missing"
        self.blocked_by_robots = False


class _Http:
    def __init__(self, ok: bool) -> None:
        self._ok = ok

    async def head(self, url: str):
        return _Probe(self._ok, 200 if self._ok else 404)


@pytest.mark.asyncio
async def test_live_url_probe_accepts_only_a_reachable_official_url(tmp_config) -> None:
    tmp_config.fixture_mode = False
    job = make_job()
    failed = PipelineState(config=tmp_config, jobs=[job], resources={"http": _Http(False)})
    await run_url_verification(failed)
    assert failed.jobs == []
    assert rejection_code(failed.rejected[0]) == "URL_VERIFICATION_FAILED"
    succeeded = PipelineState(config=tmp_config, jobs=[make_job()], resources={"http": _Http(True)})
    await run_url_verification(succeeded)
    assert succeeded.jobs[0].url_verified is True
    assert "greenhouse.io" in succeeded.jobs[0].direct_application_url
