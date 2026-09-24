from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.models.job import AppliedFlag, ApplicationStatus
from src.services.xlsx import ALL_COLUMNS, BASE_COLUMNS, sanitize_cell, write_workbooks
from tests.conftest import make_job


def test_column_order() -> None:
    assert BASE_COLUMNS[:16] == (
        "Company",
        "Job Title",
        "Normalized Role",
        "Location",
        "Remote Type",
        "Employment Type",
        "Posted Date",
        "Updated Date",
        "Job ID",
        "Source",
        "Direct Application URL",
        "Applied",
        "Status",
        "Found At",
        "Visa Sponsorship",
        "Sponsorship Evidence",
    )


def test_workbook_quality_and_hyperlinks(tmp_path: Path) -> None:
    job = make_job()
    current = tmp_path / "jobs.xlsx"
    write_workbooks([job], current_path=current, archive_dir=tmp_path / "archive", write_archive=False)
    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    assert headers[: len(ALL_COLUMNS)] == list(ALL_COLUMNS)
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref
    url_cell = sheet.cell(2, headers.index("Direct Application URL") + 1)
    assert url_cell.hyperlink is not None
    assert "greenhouse.io" in url_cell.hyperlink.target
    applied_cell = sheet.cell(2, headers.index("Applied") + 1)
    assert applied_cell.value == AppliedFlag.NOT_APPLIED.value
    assert any("L" in str(dv.sqref) for dv in sheet.data_validations.dataValidation)
    wb.close()


def test_formula_injection_protection() -> None:
    assert sanitize_cell("=HYPERLINK(\"http://evil\")") == "'=HYPERLINK(\"http://evil\")"
    assert sanitize_cell("=SUM(A1:A2)") == "'=SUM(A1:A2)"
    assert sanitize_cell("+cmd") == "'+cmd"
    assert sanitize_cell("+123") == "'+123"
    assert sanitize_cell("-2") == "'-2"
    assert sanitize_cell("-123") == "'-123"
    assert sanitize_cell("@SUM(A1)") == "'@SUM(A1)"
    assert sanitize_cell("@example") == "'@example"
    assert sanitize_cell("Software Engineer") == "Software Engineer"


def test_manual_values_and_same_day_rerun(tmp_path: Path) -> None:
    job = make_job(job_id="keep-me")
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks([job], current_path=current, archive_dir=archive, write_archive=True)

    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("Applied") + 1, AppliedFlag.APPLIED.value)
    sheet.cell(2, headers.index("Status") + 1, ApplicationStatus.INTERVIEW.value)
    wb.save(current)
    wb.close()

    from src.services.xlsx import iter_workbook_rows

    existing = list(iter_workbook_rows(current))
    rerun = make_job(job_id="keep-me", job_title="Software Engineer (repost)")
    write_workbooks(
        [rerun],
        current_path=current,
        archive_dir=archive,
        existing_rows=existing,
        write_archive=True,
    )

    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("Applied") + 1).value == AppliedFlag.APPLIED.value
    assert sheet.cell(2, headers.index("Status") + 1).value == ApplicationStatus.INTERVIEW.value
    assert sheet.max_row == 2
    wb.close()


def test_sponsorship_columns_are_informational(tmp_path: Path) -> None:
    from src.models.job import SponsorshipEvidence, SponsorshipScope, VisaSponsorshipStatus

    job = make_job()
    job.apply_sponsorship(
        SponsorshipEvidence(
            status=VisaSponsorshipStatus.NOT_SUPPORTED,
            evidence="Current job posting states that visa sponsorship is unavailable.",
            scope=SponsorshipScope.JOB_SPECIFIC,
        )
    )
    current = tmp_path / "jobs.xlsx"
    write_workbooks([job], current_path=current, archive_dir=tmp_path / "a", write_archive=False)
    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("Visa Sponsorship") + 1).value == "Not Supported"
    assert "unavailable" in (sheet.cell(2, headers.index("Sponsorship Evidence") + 1).value or "")
    wb.close()


def _headers(sheet):
    return [cell.value for cell in sheet[1]]


def _col(headers, name: str) -> int:
    return headers.index(name) + 1


def test_reordered_rows_keep_identity(tmp_path: Path) -> None:
    job_a = make_job(job_id="A", job_title="Software Engineer")
    job_b = make_job(job_id="B", job_title="Software Engineer")
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks([job_a, job_b], current_path=current, archive_dir=archive, write_archive=False)

    wb = load_workbook(current)
    sheet = wb.active
    headers = _headers(sheet)
    sheet.cell(2, _col(headers, "Applied"), AppliedFlag.APPLIED.value)
    sheet.cell(2, _col(headers, "Status"), ApplicationStatus.APPLIED.value)
    sheet.cell(3, _col(headers, "Applied"), AppliedFlag.NOT_APPLIED.value)
    sheet.cell(3, _col(headers, "Status"), ApplicationStatus.INTERVIEW.value)
    wb.save(current)
    wb.close()

    from src.services.xlsx import iter_workbook_rows

    existing = list(iter_workbook_rows(current))
    job_c = make_job(job_id="C", job_title="Software Engineer")
    write_workbooks(
        [
            make_job(job_id="B", job_title="Software Engineer"),
            make_job(job_id="A", job_title="Software Engineer"),
            job_c,
        ],
        current_path=current,
        archive_dir=archive,
        existing_rows=existing,
        write_archive=False,
    )

    wb = load_workbook(current)
    sheet = wb.active
    headers = _headers(sheet)
    rows = {
        sheet.cell(r, _col(headers, "Job ID")).value: (
            sheet.cell(r, _col(headers, "Applied")).value,
            sheet.cell(r, _col(headers, "Status")).value,
        )
        for r in range(2, sheet.max_row + 1)
    }
    assert rows["B"] == (AppliedFlag.NOT_APPLIED.value, ApplicationStatus.INTERVIEW.value)
    assert rows["A"] == (AppliedFlag.APPLIED.value, ApplicationStatus.APPLIED.value)
    assert rows["C"] == (AppliedFlag.NOT_APPLIED.value, ApplicationStatus.NOT_STARTED.value)
    assert [sheet.cell(r, _col(headers, "Job ID")).value for r in range(2, 5)] == ["B", "A", "C"]
    wb.close()


def test_same_title_different_job_ids_are_separate_rows(tmp_path: Path) -> None:
    current = tmp_path / "jobs.xlsx"
    write_workbooks(
        [
            make_job(job_id="100", job_title="Software Engineer"),
            make_job(job_id="200", job_title="Software Engineer"),
        ],
        current_path=current,
        archive_dir=tmp_path / "archive",
        write_archive=False,
    )
    wb = load_workbook(current)
    sheet = wb.active
    headers = _headers(sheet)
    ids = [sheet.cell(r, _col(headers, "Job ID")).value for r in range(2, sheet.max_row + 1)]
    assert ids == ["100", "200"]
    wb.close()


def test_successful_zero_jobs_writes_empty_snapshot(tmp_path: Path) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks([], current_path=current, archive_dir=archive, write_archive=True)
    wb = load_workbook(current)
    assert wb.active.max_row == 1
    wb.close()
    archives = list(archive.glob("*.xlsx"))
    assert len(archives) == 1
    archived = load_workbook(archives[0])
    assert archived.active.max_row == 1
    archived.close()


def test_disappeared_job_stays_in_older_archive(tmp_path: Path) -> None:
    from datetime import timedelta

    from src.utils.dates import utcnow

    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    job_a = make_job(job_id="A", job_title="Software Engineer")
    write_workbooks([job_a], current_path=current, archive_dir=archive, write_archive=True)
    today = archive / f"{utcnow().strftime('%Y-%m-%d')}.xlsx"
    yesterday = archive / f"{(utcnow() - timedelta(days=1)).strftime('%Y-%m-%d')}.xlsx"
    today.rename(yesterday)

    write_workbooks(
        [make_job(job_id="B")],
        current_path=current,
        archive_dir=archive,
        write_archive=True,
    )

    old = load_workbook(yesterday)
    old_headers = _headers(old.active)
    assert old.active.cell(2, _col(old_headers, "Job ID")).value == "A"
    old.close()

    wb = load_workbook(current)
    headers = _headers(wb.active)
    ids = [wb.active.cell(r, _col(headers, "Job ID")).value for r in range(2, wb.active.max_row + 1)]
    assert ids == ["B"]
    wb.close()


def test_fatal_write_preserves_previous_workbook(tmp_path: Path, monkeypatch) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    current.parent.mkdir(parents=True, exist_ok=True)
    current.write_bytes(b"known-good-workbook")

    def boom(self, path) -> None:
        raise OSError("xlsx generation failed")

    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", boom)
    with pytest.raises(OSError):
        write_workbooks(
            [make_job()],
            current_path=current,
            archive_dir=archive,
            write_archive=True,
        )

    assert current.read_bytes() == b"known-good-workbook"
    assert list(archive.glob("*.xlsx")) == []
    assert list(tmp_path.rglob("*.tmp")) == []
