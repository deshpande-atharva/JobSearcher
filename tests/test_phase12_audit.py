"""Phase 12: fresh-job audit, source contract, and qualification regressions."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest

from src.agents.dedup_agent import run_dedup
from src.agents.discovery_agent import _dedupe_raw
from src.agents.freshness_agent import run_freshness
from src.agents.h1b_sponsorship_agent import run_h1b_enrichment
from src.agents.location_agent import run_location_employment
from src.agents.qc_agent import run_quality_control
from src.agents.role_agent import run_role_classification
from src.agents.seniority_agent import run_seniority
from src.agents.url_agent import run_url_verification
from src.graph.pipeline import build_graph
from src.llm.base import NullLLMProvider
from src.models.config import CompanyConfig
from src.models.job import AppliedFlag, DateSource, RawJobPosting, RejectionReason
from src.models.state import PipelineState, SourceHealth
from src.services.fresh_preview import build_fresh_preview
from src.services.freshness import is_fresh
from src.services.production_report import source_contract
from src.services.roles import classify_role
from src.services.xlsx import iter_workbook_rows, write_workbooks
from src.sources.base import FetchResult, HttpClient, SourceContext
from src.sources.workday import WorkdaySource, cxs_request_body
from src.utils.logging import get_logger
from tests.conftest import make_job
from tests.test_llm_reliability import _ask, _provider

NOW = datetime(2026, 9, 27, 12, tzinfo=timezone.utc)
ADOBE = "https://adobe.wd5.myworkdayjobs.com/en-US/external_experienced"


class _Response:
    def __init__(self, status: int = 200, text: str = "ok") -> None:
        self.status_code = status
        self.headers: dict[str, str] = {}
        self.text = text
        self.url = "https://example.com/jobs"
        self.is_success = 200 <= status < 300


class _Client:
    def __init__(self, status: int = 200, text: str = "ok") -> None:
        self.calls = 0
        self.status = status
        self.text = text

    async def request(self, method: str, url: str, **kwargs: object) -> _Response:
        self.calls += 1
        return _Response(self.status, self.text)

    async def aclose(self) -> None:
        return None


def _http(config, client: _Client, *, retries: int = 0) -> HttpClient:
    http = HttpClient(config)
    http._client = client
    http._respect_robots = False
    http._retries = retries
    http._throttle._min_delay = 0
    return http


def test_production_graph_is_unchanged() -> None:
    edges = {(edge.source, edge.target) for edge in build_graph().get_graph().edges}
    assert ("discovery", "extraction") in edges
    assert ("extraction", "role") in edges
    assert ("role", "seniority") in edges
    assert ("seniority", "location") in edges
    assert ("location", "freshness") in edges
    assert ("freshness", "url") in edges
    assert ("url", "h1b") in edges
    assert ("h1b", "dedup") in edges
    assert ("dedup", "qc") in edges
    assert ("qc", "intelligence") in edges


def test_role_matrix_documents_ambiguous_families(tmp_config) -> None:
    roles = tmp_config.roles
    must_pass = (
        "Software Engineer",
        "Software Engineer II",
        "Software Developer",
        "Backend Engineer",
        "Backend Software Engineer",
        "Frontend Engineer",
        "Full Stack Engineer",
        "Full-Stack Software Engineer",
        "Platform Engineer",
        "Infrastructure Engineer",
        "Systems Software Engineer",
        "Product Engineer",
        "AI Software Engineer",
        "ML Platform Engineer",
        "Junior Software Engineer",
        "Junior Backend Engineer",
        "Junior Platform Engineer",
    )
    for title in must_pass:
        verdict = classify_role(title, "Write production code.", roles)
        assert verdict.is_software_engineering is True, title
        assert verdict.needs_llm is False, title
    must_fail = (
        "Account Manager",
        "Enterprise Account Manager",
        "Marketing Manager",
        "Product Manager",
        "Program Manager",
        "Project Manager",
        "Counsel",
        "Privacy Counsel",
        "Legislative Counsel",
        "Business Development Manager",
        "Sales Engineer",
        "Solutions Architect",
        "Engineering Manager",
        "Data Scientist",
        "Research Scientist",
        "Junior Account Manager",
    )
    for title in must_fail:
        verdict = classify_role(title, "Own the account plan.", roles)
        assert verdict.is_software_engineering is False, title
        assert verdict.needs_llm is False, title

    # Bare technical titles stay unresolved. They are not auto-accepted.
    for title in ("Engineer", "Technical Engineer", "Developer Advocate"):
        verdict = classify_role(title, "", roles)
        assert verdict.is_software_engineering is False, title
        assert verdict.needs_llm is True, title
    # Management is decided before any software keyword.
    for title in ("Software Engineering Manager", "Engineering Program Manager"):
        verdict = classify_role(title, "Write production code and code review.", roles)
        assert verdict.is_software_engineering is False, title
        assert verdict.needs_llm is False, title
    # solutions engineer is a configured core product-engineer keyword.
    solutions = classify_role("Solutions Engineer", "", roles)
    assert solutions.is_software_engineering is True
    assert solutions.family == "PRODUCT_ENGINEER"
    assert solutions.needs_llm is False
    # Data and ML titles stay ambiguous unless the description shows software work.
    data = classify_role("Data Engineer", "Analyze datasets.", roles)
    assert data.needs_llm is True and data.is_software_engineering is False
    ml = classify_role("ML Engineer", "", roles)
    assert ml.needs_llm is True and ml.is_software_engineering is False
    ai_open = classify_role("AI Engineer", "", roles)
    assert ai_open.needs_llm is True and ai_open.is_software_engineering is False
    ai_software = classify_role(
        "AI Engineer",
        "Write production code and run code review for backend services.",
        roles,
    )
    assert ai_software.is_software_engineering is True
    assert ai_software.needs_llm is False
    scientist = classify_role("Data Scientist", "Write production code and code review.", roles)
    assert scientist.is_software_engineering is False
    assert scientist.needs_llm is False


def test_freshness_boundaries_and_duplicate_timestamps() -> None:
    posted = RawJobPosting(
        source="greenhouse",
        company_name="Example",
        title="Software Engineer",
        job_id="1",
        apply_url="https://boards.greenhouse.io/example/jobs/1",
        posted_at=NOW - timedelta(hours=24),
        date_source=DateSource.POSTED_DATE,
        discovered_at=NOW,
    )
    assert is_fresh(posted, 24, now=NOW)[0] is True
    assert is_fresh(posted.model_copy(update={"posted_at": NOW - timedelta(hours=24, seconds=1)}), 24, now=NOW)[0] is False
    missing = RawJobPosting(source="greenhouse", company_name="Example", discovered_at=NOW)
    assert is_fresh(missing, 24, now=NOW) == (False, None)
    discovered_only = missing.model_copy(update={"date_source": DateSource.DISCOVERED_DATE})
    assert is_fresh(discovered_only, 24, now=NOW) == (False, None)
    fresh = posted.model_copy(update={"posted_at": NOW - timedelta(hours=1)})
    stale = fresh.model_copy(update={"posted_at": NOW - timedelta(days=10), "source": "lever"})
    assert _dedupe_raw([fresh, stale])[0].posted_at == fresh.posted_at
    assert _dedupe_raw([stale, fresh])[0].posted_at == fresh.posted_at
    unknown = fresh.model_copy(update={"job_id": "2", "apply_url": "https://boards.greenhouse.io/example/jobs/2", "posted_at": None, "date_source": DateSource.UNKNOWN})
    later = unknown.model_copy(update={"posted_at": NOW - timedelta(hours=2), "date_source": DateSource.POSTED_DATE, "source": "ashby"})
    adopted = _dedupe_raw([unknown, later])[0]
    assert adopted.posted_at == later.posted_at
    assert adopted.discovered_at == unknown.discovered_at


def test_source_contract_covers_success_empty_partial_and_failed() -> None:
    success = SourceHealth(source="greenhouse", attempted=2, successful=2, jobs_discovered=4)
    assert source_contract("greenhouse", success, {"by_source": {"greenhouse": {"caps": 0}}})["status"] == "SUCCESS"
    empty = SourceHealth(source="lever", attempted=1, successful=1, jobs_discovered=0)
    assert source_contract("lever", empty, {})["status"] == "EMPTY"
    partial = SourceHealth(source="workday", attempted=2, successful=1, failed=1, jobs_discovered=3)
    assert source_contract("workday", partial, {})["status"] == "PARTIAL"
    capped = SourceHealth(source="workday", attempted=1, successful=1, jobs_discovered=2000)
    capped_row = source_contract("workday", capped, {"by_source": {"workday": {"caps": 1}}})
    assert capped_row["status"] == "PARTIAL"
    assert capped_row["incomplete_boards"] == 1
    failed = SourceHealth(source="icims", attempted=2, failed=2)
    row = source_contract("icims", failed, {}, failed_companies=["Acme", "Other"])
    assert row["status"] == "FAILED"
    assert row["failed_companies"] == ["Acme", "Other"]


def test_cxs_body_rejects_extra_filters() -> None:
    body = cxs_request_body(limit=20, offset=0, search_text="software engineer")
    assert list(body) == ["appliedFacets", "limit", "offset", "searchText"]
    assert body["appliedFacets"] == {}
    assert body["limit"] == 20


def test_archive_replacement_leaves_previous_dates_unchanged(tmp_path) -> None:
    archive = tmp_path / "archive"
    archive.mkdir()
    previous = archive / "2026-09-24.xlsx"
    previous.write_bytes(b"historical-archive")
    digest = hashlib.sha256(previous.read_bytes()).hexdigest()
    current, written = write_workbooks([], current_path=tmp_path / "current" / "jobs.xlsx", archive_dir=archive)
    assert current.exists()
    assert written is not None and written.exists()
    assert previous.read_bytes() == b"historical-archive"
    assert hashlib.sha256(previous.read_bytes()).hexdigest() == digest
    assert written.name != previous.name


def test_applied_status_survives_and_distinct_ids_stay_separate(tmp_path) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks(
        [
            make_job(job_id="123", job_title="Software Engineer"),
            make_job(job_id="456", job_title="Software Engineer"),
        ],
        current_path=current,
        archive_dir=archive,
        write_archive=False,
    )
    from openpyxl import load_workbook

    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("Applied") + 1, AppliedFlag.APPLIED.value)
    sheet.cell(2, headers.index("Status") + 1, "Applied")
    sheet.cell(3, headers.index("Applied") + 1, AppliedFlag.NOT_APPLIED.value)
    sheet.cell(3, headers.index("Status") + 1, "Interested")
    wb.save(current)
    wb.close()
    write_workbooks(
        [
            make_job(job_id="123", job_title="Software Engineer"),
            make_job(job_id="456", job_title="Software Engineer"),
        ],
        current_path=current,
        archive_dir=archive,
        existing_rows=list(iter_workbook_rows(current)),
        write_archive=False,
    )
    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    rows = {
        str(sheet.cell(row, headers.index("Job ID") + 1).value): (
            sheet.cell(row, headers.index("Applied") + 1).value,
            sheet.cell(row, headers.index("Status") + 1).value,
        )
        for row in range(2, sheet.max_row + 1)
    }
    assert rows["123"][0] == AppliedFlag.APPLIED.value
    assert rows["123"][1] == "Applied"
    assert rows["456"][0] == AppliedFlag.NOT_APPLIED.value
    assert rows["456"][1] == "Interested"
    assert len(rows) == 2


@pytest.mark.asyncio
async def test_cache_keys_do_not_mix_range_or_accept(tmp_config) -> None:
    client = _Client()
    http = _http(tmp_config, client)
    url = "https://example.com/jobs"
    await http.get_text(url)
    await http.get_text(url)
    assert client.calls == 1
    await http.request("GET", url, headers={"Range": "bytes=0-10"})
    await http.request("GET", url, headers={"Range": "bytes=100-200"})
    assert client.calls == 3
    await http.request("GET", url, expect_json=True)
    assert client.calls == 4
    await http.request("POST", url, json_body={"offset": 0})
    await http.request("POST", url, json_body={"offset": 0})
    assert client.calls == 6
    await http.aclose()

    missing = _Client(status=404, text="")
    http404 = _http(tmp_config, missing, retries=3)
    await http404.get_text(url)
    assert missing.calls == 1
    await http404.aclose()

    broken = _Client(status=503, text="")
    http503 = _http(tmp_config, broken, retries=0)
    await http503.get_text("https://example.com/other")
    await http503.get_text("https://example.com/other")
    assert broken.calls == 2
    await http503.aclose()


@pytest.mark.asyncio
async def test_circuit_failed_probe_stays_open_and_a_new_provider_starts_closed() -> None:
    provider = _provider(
        [ValueError("down"), ValueError("down"), ValueError("still down")],
        circuit_breaker_failures=2,
        circuit_reset_seconds=0,
        retry_attempts=0,
    )
    assert await _ask(provider) is None
    assert await _ask(provider) is None
    assert provider.circuit_state == "OPEN"
    assert await _ask(provider) is None
    assert provider.circuit_state == "OPEN"
    assert provider.stats.recovery_attempts == 1
    assert provider.stats.skipped_circuit_open == 0
    fresh = _provider(['{"value": "ok"}'])
    assert fresh.circuit_state == "CLOSED"
    assert (await _ask(fresh)).value == "ok"


@pytest.mark.asyncio
async def test_qualifying_fresh_jobs_reach_final_and_controls_do_not(tmp_config) -> None:
    visa = tmp_config.settings.visa.model_copy(update={"enabled": False})
    tmp_config.settings = tmp_config.settings.model_copy(update={"visa": visa})

    async def qualify(job):
        state = PipelineState(config=tmp_config, jobs=[job], resources={"llm": NullLLMProvider()})
        await run_role_classification(state)
        await run_seniority(state)
        await run_location_employment(state)
        await run_freshness(state)
        await run_url_verification(state)
        await run_h1b_enrichment(state)
        await run_dedup(state)
        await run_quality_control(state)
        return state

    swe = make_job(
        job_title="Software Engineer",
        location="Seattle, WA",
        description="Required Qualifications\n0-2 years of experience.\nWrite production code.",
        posted_at=NOW - timedelta(hours=2),
        date_source=DateSource.POSTED_DATE,
    )
    junior = make_job(
        job_id="junior",
        job_title="Junior Backend Engineer",
        location="Remote - United States",
        description="Required Qualifications\n1+ years of experience.\nWrite production code.",
        posted_at=NOW - timedelta(hours=3),
        date_source=DateSource.POSTED_DATE,
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/junior",
    )
    senior_low = make_job(
        job_id="senior-low",
        job_title="Senior Software Engineer",
        location="Boston, MA",
        description="Required Qualifications\n0-2 years of experience.\nWrite production code.",
        posted_at=NOW - timedelta(hours=2),
        date_source=DateSource.POSTED_DATE,
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/senior-low",
    )
    for job in (swe, junior, senior_low):
        state = await qualify(job)
        assert [item.job_id for item in state.jobs] == [job.job_id]
        assert state.summary.jobs_accepted == 1

    bare = make_job(
        job_id="bare",
        location="Remote",
        description="Required Qualifications\n0-2 years of experience.\nWrite production code.",
        posted_at=NOW - timedelta(hours=2),
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/bare",
    )
    bare_state = await qualify(bare)
    assert bare_state.jobs == []
    assert bare_state.rejected[-1].reason is RejectionReason.LOCATION

    canada = make_job(
        job_id="canada",
        location="Remote - Canada",
        description="Required Qualifications\n0-2 years of experience.\nWrite production code.",
        posted_at=NOW - timedelta(hours=2),
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/canada",
    )
    canada_state = await qualify(canada)
    assert canada_state.jobs == []
    assert canada_state.rejected[-1].reason is RejectionReason.LOCATION

    senior_high = make_job(
        job_id="senior-high",
        job_title="Senior Software Engineer",
        location="Boston, MA",
        description="Required Qualifications\n5+ years of experience.",
        posted_at=NOW - timedelta(hours=2),
        direct_application_url="https://boards.greenhouse.io/acmerobotics/jobs/senior-high",
    )
    high_state = await qualify(senior_high)
    assert high_state.jobs == []
    assert high_state.rejected[-1].reason is RejectionReason.SENIORITY


@pytest.mark.asyncio
async def test_h1b_lookup_does_not_run_after_a_freshness_rejection(tmp_config) -> None:
    class Spy:
        def __init__(self) -> None:
            self.calls = 0

        async def lookup(self, company: str, aliases=()):
            self.calls += 1
            return None

    stale = make_job(posted_at=NOW - timedelta(days=40), date_source=DateSource.POSTED_DATE)
    state = PipelineState(config=tmp_config, jobs=[stale], resources={"h1b": Spy(), "llm": NullLLMProvider()})
    await run_freshness(state)
    assert state.jobs == []
    await run_h1b_enrichment(state)
    assert state.resources["h1b"].calls == 0
    assert state.rejected[-1].reason is RejectionReason.FRESHNESS


@pytest.mark.asyncio
async def test_intelligence_restores_identity_if_a_fit_step_mutates_it(tmp_config, monkeypatch) -> None:
    import src.agents.intelligence_agent as intel

    job = make_job()
    original = (job.job_id, job.direct_application_url, job.posted_at, job.date_source, job.source)

    def evil(posting, profile, roles):
        job.job_id = "changed"
        job.direct_application_url = "https://jobright.ai/jobs/1"
        job.posted_at = None
        job.source = "jobright"
        return SimpleNamespace()

    monkeypatch.setattr(intel, "load_candidate_profile", lambda path: SimpleNamespace(profile_version=None))
    monkeypatch.setattr(intel, "merge_resume_profile", lambda search, path: SimpleNamespace(profile_version=1))
    monkeypatch.setattr(intel, "evaluate_job", evil)
    monkeypatch.setattr(
        intel,
        "review_fit",
        lambda fit, posting, profile: SimpleNamespace(unsupported_claims=["invented skill"]),
    )
    state = PipelineState(config=tmp_config, jobs=[job])
    await intel.run_job_intelligence(state)
    assert (job.job_id, job.direct_application_url, job.posted_at, job.date_source, job.source) == original
    assert state.jobs == [job]
    assert state.summary.unsupported_claims_removed == 1


def test_fresh_audit_counters_follow_the_pipeline_reason(tmp_config) -> None:
    posting = RawJobPosting(
        source="greenhouse",
        company_name="Airbnb",
        title="Legislative Counsel",
        job_id="1",
        apply_url="https://boards.greenhouse.io/airbnb/jobs/1",
        location_raw="Remote",
        posted_at=datetime.now(timezone.utc) - timedelta(hours=2),
        date_source=DateSource.POSTED_DATE,
    )
    state = PipelineState(config=tmp_config, raw_postings=[posting])
    state.reject(make_job(job_title="Legislative Counsel", direct_application_url=posting.apply_url), RejectionReason.ROLE, "no software-engineering title")
    audit = build_fresh_preview(state)
    assert audit["counters"]["fresh_authoritative_total"] == 1
    assert audit["counters"]["fresh_role_reject"] == 1
    assert audit["counters"]["fresh_final_candidates"] == 0
    assert audit["jobs"][0]["official_url"].startswith("https://boards.greenhouse.io/")
    assert "description" not in audit["jobs"][0]


@pytest.mark.asyncio
async def test_workday_fresh_recall_distinguishes_visible_recovered_and_incomplete(tmp_config) -> None:
    class Board:
        def __init__(self, mode: str) -> None:
            self.mode = mode
            self.bodies: list[dict] = []

        async def request(self, method: str, url: str, **kwargs):
            body = dict(kwargs.get("json_body") or {})
            self.bodies.append(body)
            text = str(body.get("searchText") or "")
            offset = int(body.get("offset") or 0)
            if text:
                if self.mode == "recovered":
                    jobs = [{"title": "Software Engineer", "externalPath": "/job/hidden", "id": "HIDDEN", "postedOn": "Posted 2 Hours Ago"}]
                elif self.mode == "visible":
                    jobs = [{"title": "Software Engineer", "externalPath": "/job/visible", "id": "VISIBLE", "postedOn": "Posted 2 Hours Ago"}]
                else:
                    jobs = [{"title": "Software Engineer", "externalPath": "/job/old", "id": "OUTSIDE-NOT-RETURNED", "postedOn": "Posted 30+ Days Ago"}]
                return FetchResult(url=url, status=200, text=json.dumps({"total": 1, "jobPostings": jobs if offset == 0 else []}))
            if offset >= 40:
                return FetchResult(url=url, status=200, text='{"total":0,"jobPostings":[]}')
            jobs = []
            for index in range(20):
                job_id = "VISIBLE" if self.mode == "visible" and offset == 0 and index == 0 else f"OLD{offset}-{index}"
                posted = "Posted 2 Hours Ago" if job_id == "VISIBLE" else "Posted 30+ Days Ago"
                jobs.append({"title": "Software Engineer", "externalPath": f"/job/{job_id}", "id": job_id, "postedOn": posted})
            total = 2000 if offset == 0 else 0
            return FetchResult(url=url, status=200, text=json.dumps({"total": total, "jobPostings": jobs}))

    async def run(mode: str):
        workday = tmp_config.settings.discovery.sources.workday.model_copy(
            update={
                "partitions_enabled": True,
                "max_partitions_per_company": 1,
                "max_jobs": 40,
                "max_jobs_per_partition": 20,
                "partition_search_texts": ["software engineer"],
            }
        )
        sources = tmp_config.settings.discovery.sources.model_copy(update={"workday": workday})
        discovery = tmp_config.settings.discovery.model_copy(update={"sources": sources})
        tmp_config.settings = tmp_config.settings.model_copy(update={"discovery": discovery})
        tmp_config.fixture_mode = False
        http = Board(mode)
        result = await WorkdaySource(
            SourceContext(config=tmp_config, http=http, logger=get_logger("phase12"))
        ).discover_result(CompanyConfig(name="NVIDIA", ats_type="workday", ats_identifier=ADOBE, careers_url=ADOBE))
        assert all(set(body) == {"appliedFacets", "limit", "offset", "searchText"} for body in http.bodies)
        assert all(body["appliedFacets"] == {} for body in http.bodies)
        return result

    visible = await run("visible")
    assert visible.diagnostics["estimated_incomplete"] is True
    assert int(visible.diagnostics["unfiltered_fresh_jobs"]) == 1
    assert int(visible.diagnostics["recovered_fresh_jobs"]) == 0
    recovered = await run("recovered")
    assert int(recovered.diagnostics["recovered_fresh_jobs"]) == 1
    assert any(job.job_id == "HIDDEN" for job in recovered.jobs)
    incomplete = await run("outside")
    assert incomplete.diagnostics["estimated_incomplete"] is True
    assert int(incomplete.diagnostics["recovered_fresh_jobs"]) == 0
    assert int(incomplete.diagnostics["unfiltered_fresh_jobs"]) == 0

