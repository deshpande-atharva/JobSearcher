"""Production daily-path hardening: output, health, caps, and configuration."""

from __future__ import annotations

from datetime import timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.llm.base import classify_provider_error
from src.models.job import AppliedFlag, ApplicationStatus, DateSource, RawJobPosting
from src.models.state import SourceHealth
from src.services.board_priority import cap_board_postings
from src.services.history import load_history
from src.services.notifications import build_email_body
from src.services.production_report import production_source_state
from src.services.xlsx import write_workbooks
from src.sources import COMPANY_SOURCES
from src.utils.dates import utcnow
from tests.conftest import make_job


class _Boom(Exception):
    def __init__(self, status: int) -> None:
        super().__init__(f"status {status}")
        self.status_code = status


def test_source_states_are_deterministic() -> None:
    assert production_source_state(SourceHealth(source="lever", attempted=1, successful=1, jobs_discovered=4)) == "SUCCESS"
    assert production_source_state(SourceHealth(source="ashby", attempted=1, successful=1, empty=1)) == "EMPTY"
    assert production_source_state(
        SourceHealth(source="greenhouse", attempted=3, successful=2, failed=1, jobs_discovered=10)
    ) == "PARTIAL"
    assert production_source_state(SourceHealth(source="workday", attempted=2, failed=2, error=2)) == "FAILED"
    assert production_source_state(SourceHealth(source="jobright")) == "DISABLED"


def test_production_configuration_invariants(tmp_config) -> None:
    sources = tmp_config.settings.discovery.sources
    assert sources.jobright.enabled is False
    assert sources.workday.browser_enabled is False
    assert sources.lever.enabled is True
    assert sources.ashby.enabled is True
    assert sources.lever.prioritize_fresh_targets is False
    assert sources.ashby.prioritize_fresh_targets is False
    assert tmp_config.freshness_hours == 24
    names = {cls.name for cls in COMPANY_SOURCES}
    assert {"greenhouse", "workday", "lever", "ashby"} <= names
    assert not Path("src/agents/lever_job_intelligence_agent.py").exists()
    assert not Path("src/agents/ashby_job_intelligence_agent.py").exists()
    workflow = Path(".github/workflows/daily_jobs.yml").read_text(encoding="utf-8")
    assert 'cron: "0 12 * * *"' in workflow
    assert 'python-version: "3.12"' in workflow
    assert "python -m src.main" in workflow
    assert "--smoke" not in workflow
    assert "--fixture" not in workflow
    assert "git add data/current data/archive" in workflow
    assert "contents: write" in workflow


def test_timestamp_order_fills_a_binding_cap_and_leaves_a_short_board_alone(tmp_config) -> None:
    older = RawJobPosting(
        source="lever",
        company_name="Example",
        title="Software Engineer",
        job_id="old",
        apply_url="https://jobs.lever.co/example/aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeee01",
        posted_at=utcnow() - timedelta(days=3),
        date_source=DateSource.POSTED_DATE,
    )
    newer = older.model_copy(
        update={
            "job_id": "new",
            "posted_at": utcnow() - timedelta(hours=2),
            "apply_url": "https://jobs.lever.co/example/aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeee02",
        }
    )
    undated = older.model_copy(
        update={
            "job_id": "none",
            "posted_at": None,
            "date_source": DateSource.UNKNOWN,
            "apply_url": "https://jobs.lever.co/example/aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeee03",
        }
    )
    settings = tmp_config.settings.discovery.sources.lever.model_copy(update={"max_jobs": 1})
    kept, capped, note = cap_board_postings(
        [older, undated, newer], tmp_config, limit=settings.max_jobs, prioritize=False
    )
    assert capped is True
    assert kept[0].job_id == "new"
    assert "timestamp_order=authoritative" in note
    untouched, open_cap, open_note = cap_board_postings(
        [older, newer], tmp_config, limit=2, prioritize=False
    )
    assert open_cap is False
    assert [item.job_id for item in untouched] == ["old", "new"]
    assert open_note == ""


def test_two_identical_workbook_runs_do_not_duplicate_rows(tmp_path: Path) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    jobs = [
        make_job(job_id="123", job_title="Software Engineer"),
        make_job(job_id="456", job_title="Backend Engineer"),
    ]
    write_workbooks(jobs, current_path=current, archive_dir=archive, write_archive=True)
    from src.services.xlsx import iter_workbook_rows

    existing = list(iter_workbook_rows(current))
    write_workbooks(jobs, current_path=current, archive_dir=archive, existing_rows=existing, write_archive=True)
    wb = load_workbook(current)
    ids = [wb.active.cell(row, 9).value for row in range(2, wb.active.max_row + 1)]
    wb.close()
    assert ids == ["123", "456"]
    assert len(list(archive.glob("*.xlsx"))) == 1


def test_applied_and_status_survive_and_a_new_id_stays_separate(tmp_path: Path) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks(
        [
            make_job(job_id="123", job_title="Software Engineer"),
            make_job(job_id="456", job_title="Backend Engineer"),
        ],
        current_path=current,
        archive_dir=archive,
        write_archive=False,
    )
    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    applied = headers.index("Applied") + 1
    status = headers.index("Status") + 1
    sheet.cell(2, applied, AppliedFlag.APPLIED.value)
    sheet.cell(2, status, ApplicationStatus.APPLIED.value)
    sheet.cell(3, applied, AppliedFlag.NOT_APPLIED.value)
    sheet.cell(3, status, "Interested")
    wb.save(current)
    wb.close()

    from src.services.xlsx import iter_workbook_rows

    existing = list(iter_workbook_rows(current))
    write_workbooks(
        [
            make_job(job_id="123", job_title="Software Engineer"),
            make_job(job_id="456", job_title="Backend Engineer"),
            make_job(job_id="789", job_title="Software Engineer"),
        ],
        current_path=current,
        archive_dir=archive,
        existing_rows=existing,
        write_archive=False,
    )
    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    rows = {
        sheet.cell(row, headers.index("Job ID") + 1).value: (
            sheet.cell(row, headers.index("Applied") + 1).value,
            sheet.cell(row, headers.index("Status") + 1).value,
        )
        for row in range(2, sheet.max_row + 1)
    }
    wb.close()
    assert rows["123"] == (AppliedFlag.APPLIED.value, ApplicationStatus.APPLIED.value)
    assert rows["456"] == (AppliedFlag.NOT_APPLIED.value, "Interested")
    assert rows["789"] == (AppliedFlag.NOT_APPLIED.value, ApplicationStatus.NOT_STARTED.value)
    assert len(rows) == 3


def test_archive_write_failure_keeps_the_current_workbook(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    current = tmp_path / "jobs.xlsx"
    current.write_bytes(b"known-good")
    archive = tmp_path / "archive"
    calls = {"n": 0}
    from src.services import xlsx as xlsx_module

    real = xlsx_module._write_sheet

    def flaky(path, rows) -> None:
        calls["n"] += 1
        if calls["n"] == 2:
            raise OSError("archive disk full")
        real(path, rows)

    monkeypatch.setattr(xlsx_module, "_write_sheet", flaky)
    with pytest.raises(OSError, match="archive disk full"):
        write_workbooks([make_job()], current_path=current, archive_dir=archive, write_archive=True)
    assert current.read_bytes() == b"known-good"
    assert list(archive.glob("*.xlsx")) == []


def test_current_replace_failure_keeps_the_previous_workbook(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    current = tmp_path / "jobs.xlsx"
    archive = tmp_path / "archive"
    write_workbooks([make_job(job_id="keep")], current_path=current, archive_dir=archive, write_archive=False)
    original = current.read_bytes()
    real_replace = __import__("os").replace

    def boom(src, dst) -> None:
        if Path(dst) == current:
            raise OSError("replace failed")
        real_replace(src, dst)

    monkeypatch.setattr("src.services.xlsx.os.replace", boom)
    with pytest.raises(OSError, match="replace failed"):
        write_workbooks([make_job(job_id="new")], current_path=current, archive_dir=archive, write_archive=True)
    assert current.read_bytes() == original


def test_malformed_existing_workbook_does_not_raise(tmp_path: Path, tmp_config) -> None:
    current = tmp_config.path(tmp_config.settings.output.current_workbook)
    current.parent.mkdir(parents=True, exist_ok=True)
    current.write_bytes(b"this is not a workbook")
    history = load_history(tmp_config)
    assert history.existing_rows == []


def test_email_mentions_source_health_and_a_zero_result(tmp_config) -> None:
    from src.models.state import RunSummary

    summary = RunSummary(jobs_accepted=0, freshness_hours_used=24)
    summary.record_source("greenhouse", success=False, jobs=0, error="down", status="ERROR")
    summary.record_source("lever", success=True, jobs=3, status="OK")
    summary.failed_sources.append("greenhouse")
    body = build_email_body(summary, [])
    assert "Final jobs: 0" in body
    assert "valid zero-result run" in body
    assert "greenhouse: FAILED" in body
    assert "lever: SUCCESS" in body
    assert "Failed sources: greenhouse" in body
    assert "password" not in body.lower()


def test_provider_errors_stay_in_the_existing_buckets() -> None:
    assert classify_provider_error(_Boom(429)) == "429"
    assert classify_provider_error(_Boom(503)) == "503"
    assert classify_provider_error(TimeoutError("timed out")) == "timeout"


async def test_one_malformed_lever_record_keeps_the_rest(tmp_config, monkeypatch: pytest.MonkeyPatch) -> None:
    from src.models.config import CompanyConfig
    from src.sources.base import SourceContext
    from src.sources.fixtures import FixtureStore
    from src.sources.lever import LeverSource
    from src.utils.logging import get_logger

    good = {
        "id": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee",
        "text": "Software Engineer",
        "hostedUrl": "https://jobs.lever.co/example/aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee",
        "categories": {"location": "Boston, MA", "commitment": "Full-time"},
        "createdAt": "1 hour ago",
        "descriptionPlain": "0-2 years. Full-time software engineer.",
    }
    monkeypatch.setattr(FixtureStore, "load_first", lambda self, source, names: ["broken", good])
    real = LeverSource._to_posting

    def flaky(self, entry, company, handle):
        if entry == "broken":
            raise RuntimeError("malformed posting")
        return real(self, entry, company, handle)

    monkeypatch.setattr(LeverSource, "_to_posting", flaky)
    source = LeverSource(SourceContext(config=tmp_config, http=None, logger=get_logger("phase7")))
    company = CompanyConfig(name="Example", ats_type="lever", ats_identifier="example")
    jobs = await source.discover(company)
    assert [job.job_id for job in jobs] == ["aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"]


async def test_two_fixture_pipeline_runs_keep_tracking(tmp_config) -> None:
    from src.graph.pipeline import run_pipeline
    from src.services.xlsx import iter_workbook_rows

    config = tmp_config.model_copy(update={"dry_run": False, "send_email": False})
    first = await run_pipeline(config)
    assert first.summary.email_status.startswith("skipped")
    current = config.path(config.settings.output.current_workbook)
    assert current.is_file()
    wb = load_workbook(current)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    if sheet.max_row < 2:
        wb.close()
        pytest.skip("fixture universe accepted no jobs")
    job_id = sheet.cell(2, headers.index("Job ID") + 1).value
    sheet.cell(2, headers.index("Applied") + 1, AppliedFlag.APPLIED.value)
    sheet.cell(2, headers.index("Status") + 1, ApplicationStatus.APPLIED.value)
    wb.save(current)
    wb.close()
    first_ids = [row.get("Job ID") for row in iter_workbook_rows(current)]

    second = await run_pipeline(config)
    second_ids = [row.get("Job ID") for row in iter_workbook_rows(current)]
    assert second_ids == first_ids
    assert len(second_ids) == len(set(second_ids))
    saved = {row.get("Job ID"): row for row in iter_workbook_rows(current)}
    assert saved[job_id]["Applied"] == AppliedFlag.APPLIED.value
    assert saved[job_id]["Status"] == ApplicationStatus.APPLIED.value
    assert second.summary.jobs_accepted == first.summary.jobs_accepted
    archives = list(config.path(config.settings.output.archive_dir).glob("*.xlsx"))
    assert len(archives) == 1
